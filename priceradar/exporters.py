from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.sax.saxutils import escape
from .core import brl

HEADERS=['Loja','Produto','Preço (BRL)','URL','Consultado em']

def export_file(path,offers):
    path=Path(path)
    # A failed export never replaces an existing destination.
    with NamedTemporaryFile(dir=path.parent,suffix=path.suffix,delete=False) as f: temp=Path(f.name)
    try:
        if path.suffix.lower()=='.xlsx':
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            wb=Workbook(); ws=wb.active; ws.title='Ofertas'; ws.append(HEADERS)
            for o in offers:
                ws.append([o.store,o.title,o.cents/100,o.url,o.captured])
                for col in (1,2,4,5): ws.cell(ws.max_row,col).data_type='s'
                ws.cell(ws.max_row,3).number_format='"R$" #,##0.00'
                ws.cell(ws.max_row,4).hyperlink=o.url
            for cell in ws[1]: cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='635BFF')
            for col,width in zip('ABCDE',[20,65,18,55,29]): ws.column_dimensions[col].width=width
            for row in ws.iter_rows(min_row=2):
                for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)
            ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
            ws.append(['Preços dos anúncios; frete, cupons e variações devem ser conferidos na loja.'])
            wb.save(temp)
        elif path.suffix.lower()=='.pdf':
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            styles=getSampleStyleSheet(); body=styles['BodyText']; body.fontSize=9
            story=[Paragraph('PriceRadar BR',styles['Title']),Paragraph('Comparativo de ofertas coletadas',styles['Heading2']),Paragraph('Confira modelo, variação, vendedor e frete. O menor preço se refere apenas aos resultados coletados; não garante produtos equivalentes.',body),Spacer(1,18)]
            rows=[['Loja / consulta','Produto / anúncio','Preço']]
            for o in offers:
                rows.append([Paragraph(escape(o.store)+'<br/>'+escape(o.captured[:19]),body),Paragraph(escape(o.title)+'<br/><link href="'+escape(o.url,{'"':'&quot;'})+'" color="#635BFF">Abrir anúncio</link>',body),brl(o.cents)])
            table=Table(rows,colWidths=[110,310,85],repeatRows=1)
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#635BFF')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor('#F1F3F9'),colors.white]),('LEFTPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),12)]))
            story.append(table)
            def footer(c,d): c.setFont('Helvetica',8); c.drawString(45,25,f'PriceRadar BR  |  Página {d.page}')
            SimpleDocTemplate(str(temp),pagesize=A4,rightMargin=45,leftMargin=45).build(story,onFirstPage=footer,onLaterPages=footer)
        else: raise ValueError('Formato não suportado')
        temp.replace(path)
    finally: temp.unlink(missing_ok=True)
